"""Extract text from uploaded tender documents.

Tenders arrive as PDF or DOCX, not as text someone pasted into a box, so
requiring paste is requiring the user to do the awkward part by hand.

PDF tables flatten badly (the same column-flattening seen in scanned BIS
annexes), so DOCX tables are read cell-by-cell to keep line items intact —
a tender's schedule of items is almost always a table, and that's precisely
the content the linter needs to read.
"""
from __future__ import annotations

import io
from typing import Any, Dict, List, Optional

MAX_UPLOAD_BYTES = 10 * 1024 * 1024

# Below this, a PDF's text layer is treated as absent rather than sparse, and
# OCR takes over.
MIN_TEXT_LAYER_CHARS = 50
OCR_DPI = 300
OCR_MAX_PAGES = 10

SUPPORTED_TYPES = {
    "application/pdf": "pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
    "text/plain": "txt",
}

SUPPORTED_EXTENSIONS = {".pdf": "pdf", ".docx": "docx", ".txt": "txt", ".md": "txt"}


class UnsupportedDocument(Exception):
    pass


def detect_kind(filename: Optional[str], content_type: Optional[str]) -> str:
    if content_type and content_type in SUPPORTED_TYPES:
        return SUPPORTED_TYPES[content_type]
    if filename:
        lowered = filename.lower()
        for ext, kind in SUPPORTED_EXTENSIONS.items():
            if lowered.endswith(ext):
                return kind
    raise UnsupportedDocument(
        f"Unsupported document type (filename={filename!r}, content_type={content_type!r}). "
        f"Supported: PDF, DOCX, TXT."
    )


def extract_pdf_text(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    pages = []
    for page in reader.pages:
        try:
            pages.append(page.extract_text() or "")
        except Exception:
            pages.append("")  # one unreadable page shouldn't lose the document
    return "\n".join(pages)


def ocr_pdf_text(data: bytes, max_pages: int = OCR_MAX_PAGES) -> str:
    """Read a scanned PDF by rendering pages and running OCR over them.

    Approach adapted from the team's `standards-retrieval/extraction.py`.
    Government tenders are routinely circulated as scans of signed printouts,
    which have no text layer at all — without this they are simply unreadable,
    and a tender auditor that can't read half its input isn't much of an
    auditor.

    OCR is slow and lossy, so it runs only as a fallback and is page-capped;
    the caller surfaces a warning telling the user to check the extracted text.
    """
    import fitz  # PyMuPDF
    import pytesseract
    from PIL import Image

    out = []
    with fitz.open(stream=data, filetype="pdf") as doc:
        for page in doc[:max_pages]:
            # 300 DPI: below this tesseract's accuracy on printed tender text
            # falls off sharply.
            pix = page.get_pixmap(dpi=OCR_DPI)
            image = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
            out.append(pytesseract.image_to_string(image))
    return "\n".join(out)


def pdf_page_count(data: bytes) -> int:
    import fitz

    with fitz.open(stream=data, filetype="pdf") as doc:
        return doc.page_count


def extract_docx_text(data: bytes) -> str:
    import docx

    document = docx.Document(io.BytesIO(data))
    lines = [p.text for p in document.paragraphs]

    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            # Word repeats a merged cell's text across the cells it spans.
            deduped = [c for i, c in enumerate(cells) if c and (i == 0 or c != cells[i - 1])]
            if deduped:
                lines.append(" | ".join(deduped))

    return "\n".join(lines)


def extract_text_with_meta(data: bytes, *, filename: Optional[str] = None,
                           content_type: Optional[str] = None,
                           allow_ocr: bool = True) -> Dict[str, Any]:
    """Extract text, falling back to OCR for scanned PDFs.

    Returns the text plus how it was obtained and any warnings, because OCR
    output is materially less trustworthy than a real text layer and the user
    has to be told which one they are looking at before acting on findings.
    """
    if len(data) > MAX_UPLOAD_BYTES:
        raise UnsupportedDocument(
            f"Document exceeds {MAX_UPLOAD_BYTES // (1024 * 1024)} MB limit."
        )

    kind = detect_kind(filename, content_type)
    warnings: List[str] = []
    method = kind

    if kind == "pdf":
        text = extract_pdf_text(data)
        if len(text.strip()) < MIN_TEXT_LAYER_CHARS and allow_ocr:
            try:
                ocr_text = ocr_pdf_text(data)
            except Exception as exc:
                raise UnsupportedDocument(
                    "This PDF has no readable text layer and OCR failed "
                    f"({exc}). Install Tesseract (`brew install tesseract`) or "
                    "supply a text-based PDF."
                )
            if ocr_text.strip():
                text = ocr_text
                method = "pdf-ocr"
                warnings.append(
                    "No text layer found, so this was read by OCR. OCR misreads "
                    "poor scans — check the extracted text before acting on any "
                    "finding."
                )
                pages = pdf_page_count(data)
                if pages > OCR_MAX_PAGES:
                    warnings.append(
                        f"Only the first {OCR_MAX_PAGES} of {pages} pages were "
                        f"read; findings below cover that portion only."
                    )
    elif kind == "docx":
        text = extract_docx_text(data)
    else:
        text = data.decode("utf-8", errors="replace")

    if not text.strip():
        raise UnsupportedDocument(
            "No text could be extracted from this document."
        )

    return {"text": text, "method": method, "warnings": warnings}


def extract_text(data: bytes, *, filename: Optional[str] = None,
                 content_type: Optional[str] = None) -> str:
    return extract_text_with_meta(
        data, filename=filename, content_type=content_type
    )["text"]
