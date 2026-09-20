"""Document upload tests. Fixtures are generated as real PDF/DOCX bytes rather
than mocked, so these exercise the actual parsers."""
import io

import pytest

from app.core.documents import (
    MAX_UPLOAD_BYTES,
    UnsupportedDocument,
    detect_kind,
    extract_text,
)

DOCX_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _make_docx(paragraphs, table_rows=None) -> bytes:
    import docx

    document = docx.Document()
    for p in paragraphs:
        document.add_paragraph(p)
    if table_rows:
        table = document.add_table(rows=0, cols=len(table_rows[0]))
        for row in table_rows:
            cells = table.add_row().cells
            for cell, value in zip(cells, row):
                cell.text = value
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def _make_pdf(lines) -> bytes:
    """Minimal single-page PDF with a text stream pypdf can read back."""
    content = "BT /F1 12 Tf 50 750 Td 14 TL\n"
    for line in lines:
        escaped = line.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)")
        content += f"({escaped}) Tj T*\n"
    content += "ET"
    stream = content.encode("latin-1")

    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]

    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, obj in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + obj + b"\nendobj\n"

    xref_at = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode()
    out += b"0000000000 65535 f \n"
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += (f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
            f"startxref\n{xref_at}\n%%EOF\n").encode()
    return bytes(out)


@pytest.mark.parametrize("filename,content_type,expected", [
    ("tender.pdf", "application/pdf", "pdf"),
    ("tender.docx", DOCX_TYPE, "docx"),
    ("tender.txt", "text/plain", "txt"),
    ("tender.PDF", None, "pdf"),           # extension fallback, case-insensitive
    ("spec.md", None, "txt"),
    ("noext", "application/pdf", "pdf"),   # content-type wins when present
])
def test_detect_kind(filename, content_type, expected):
    assert detect_kind(filename, content_type) == expected


def test_detect_kind_rejects_unsupported():
    with pytest.raises(UnsupportedDocument):
        detect_kind("tender.xlsx", "application/vnd.ms-excel")


def test_extract_plain_text():
    text = extract_text(b"Steel to IS 2062:2006.", filename="a.txt", content_type="text/plain")
    assert "IS 2062:2006" in text


def test_extract_docx_paragraphs():
    data = _make_docx(["Supply of steel to IS 2062:2006.", "All items ISI marked."])
    text = extract_text(data, filename="t.docx", content_type=DOCX_TYPE)
    assert "IS 2062:2006" in text
    assert "ISI marked" in text


def test_extract_docx_table_rows_kept_as_line_items():
    """A tender's schedule of items is nearly always a table -- exactly the
    content the linter has to read."""
    data = _make_docx(
        ["Schedule of items"],
        table_rows=[["1", "Structural steel, IS 2062:2006", "10 MT"],
                    ["2", "Helmets for riders", "200 nos"]],
    )
    text = extract_text(data, filename="t.docx", content_type=DOCX_TYPE)
    assert "IS 2062:2006" in text
    assert "Helmets for riders" in text
    lines = [ln for ln in text.splitlines() if "|" in ln]
    assert len(lines) == 2


def test_extract_pdf_text():
    data = _make_pdf(["Supply of structural steel to IS 2062:2006",
                      "Helmets conforming to relevant IS"])
    text = extract_text(data, filename="t.pdf", content_type="application/pdf")
    assert "IS 2062:2006" in text


def test_empty_document_rejected_with_actionable_message():
    data = _make_pdf([])
    with pytest.raises(UnsupportedDocument) as exc:
        extract_text(data, filename="t.pdf", content_type="application/pdf")
    assert "OCR" in str(exc.value)


def test_oversized_upload_rejected():
    with pytest.raises(UnsupportedDocument) as exc:
        extract_text(b"x" * (MAX_UPLOAD_BYTES + 1), filename="t.txt", content_type="text/plain")
    assert "limit" in str(exc.value).lower()


def test_lint_upload_endpoint_end_to_end(client):
    data = _make_docx(
        ["Tender for departmental supplies"],
        table_rows=[["1", "Structural steel sections to IS 2062:2006", "10 MT"]],
    )
    resp = client.post(
        "/lint/upload",
        files={"file": ("tender.docx", data, DOCX_TYPE)},
        params={"suggest_missing": "false"},
    )
    assert resp.status_code == 200
    body = resp.json()

    assert body["filename"] == "tender.docx"
    assert body["characters_extracted"] > 0
    rules = {f["rule_id"] for f in body["findings"]}
    assert "superseded_citation" in rules

    # spans must index into the returned extracted text, not the original file
    finding = next(f for f in body["findings"] if f["rule_id"] == "superseded_citation")
    start, end = finding["span"]
    assert body["document_text"][start:end] == "IS 2062:2006"


def test_lint_upload_rejects_unsupported_type(client):
    resp = client.post(
        "/lint/upload",
        files={"file": ("sheet.xlsx", b"junk", "application/vnd.ms-excel")},
    )
    assert resp.status_code == 400
